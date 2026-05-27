from __future__ import annotations

import argparse
import re
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:
    from PIL import Image, ImageDraw, ImageFont
except Exception:  # pragma: no cover - converter still works without image generation.
    Image = None
    ImageDraw = None
    ImageFont = None


INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
NUMBER_RE = re.compile(r"^\d+\.\s+(.+)$")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
      tc_mar = OxmlElement("w:tcMar")
      tc_pr.append(tc_mar)

    for margin, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_dxa))


def shade_paragraph(paragraph, fill: str = "F2F4F7") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def apply_document_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles.add_style("ReportTitle", 1)
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(31, 77, 120)
    title.paragraph_format.space_after = Pt(12)

    code = styles.add_style("CodeBlock", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(8.5)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_formatted_text(paragraph, text: str) -> None:
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(31, 77, 120)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    add_formatted_text(paragraph, text)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def table_widths(column_count: int) -> list[float]:
    patterns = {
        2: [1.8, 4.7],
        3: [1.2, 2.2, 3.1],
        4: [0.7, 1.6, 2.1, 2.1],
    }
    return patterns.get(column_count, [6.5 / column_count] * column_count)


def add_table(doc: Document, table_lines: list[str]) -> None:
    rows = [split_table_row(line) for line in table_lines]
    rows = [row for row in rows if not is_separator_row(row)]
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table.autofit = False
    set_table_width(table)

    widths = table_widths(column_count)

    for row_idx, row in enumerate(rows):
        for col_idx in range(column_count):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(widths[col_idx])
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            text = row[col_idx] if col_idx < len(row) else ""
            add_formatted_text(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if row_idx == 0:
                    run.bold = True
            if row_idx == 0:
                set_cell_shading(cell, "E8EEF5")

    doc.add_paragraph()


def load_font(size: int):
    if ImageFont is None:
        return None

    font_paths = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for path in font_paths:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def arrow(draw, start, end, fill="#334155", width=3):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        points = [(x2, y2), (x2 - 14, y2 - 7), (x2 - 14, y2 + 7)]
    else:
        points = [(x2, y2), (x2 + 14, y2 - 7), (x2 + 14, y2 + 7)]
    draw.polygon(points, fill=fill)


def poly_arrow(draw, points, fill="#334155", width=3):
    if len(points) < 2:
        return
    for start, end in zip(points[:-2], points[1:-1]):
        draw.line([start, end], fill=fill, width=width)
    arrow(draw, points[-2], points[-1], fill=fill, width=width)


def draw_box(draw, xy, title, lines, fill, outline="#334155"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=2)
    title_font = load_font(24)
    body_font = load_font(16)
    draw.text(((x1 + x2) / 2, y1 + 22), title, fill="#0f172a", font=title_font, anchor="mt")
    y = y1 + 58
    for line in lines:
        draw.text(((x1 + x2) / 2, y), line, fill="#475569", font=body_font, anchor="mt")
        y += 23


def draw_lifeline_diagram(output_path: Path, title: str, participants: list[str], messages: list[tuple[int, int, str]]) -> None:
    if Image is None:
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    width = 1600
    height = max(820, 260 + len(messages) * 76)
    img = Image.new("RGB", (width, height), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_font = load_font(36)
    participant_font = load_font(20)
    label_font = load_font(17)

    draw.text((width / 2, 34), title, fill="#0f172a", font=title_font, anchor="mt")

    left = 170
    right = width - 170
    top = 120
    bottom = height - 70
    step = (right - left) / (len(participants) - 1)
    xs = [left + i * step for i in range(len(participants))]

    for x, participant in zip(xs, participants):
        draw.rounded_rectangle((x - 110, top, x + 110, top + 58), radius=12, fill="#e0f2fe", outline="#334155", width=2)
        draw.text((x, top + 29), participant, fill="#0f172a", font=participant_font, anchor="mm")
        draw.line((x, top + 58, x, bottom), fill="#94a3b8", width=2)

    y = top + 125
    for from_idx, to_idx, label in messages:
        x1 = xs[from_idx]
        x2 = xs[to_idx]
        direction = 1 if x2 >= x1 else -1
        line_start = (x1 + direction * 12, y)
        line_end = (x2 - direction * 12, y)
        arrow(draw, line_start, line_end, fill="#334155", width=3)
        wrapped = textwrap.wrap(label, width=32)
        label_y = y - 48
        for line in wrapped[:2]:
            center_x = (x1 + x2) / 2
            text_bbox = draw.textbbox((center_x, label_y), line, font=label_font, anchor="mm")
            draw.rounded_rectangle(
                (text_bbox[0] - 7, text_bbox[1] - 3, text_bbox[2] + 7, text_bbox[3] + 3),
                radius=6,
                fill="#f8fafc",
            )
            draw.text((center_x, label_y), line, fill="#0f172a", font=label_font, anchor="mm")
            label_y += 22
        y += 76

    img.save(output_path)


def generate_auth_flow_png(output_path: Path) -> None:
    draw_lifeline_diagram(
        output_path,
        "Luồng đăng nhập / đăng ký",
        ["Người dùng", "Auth Screen", "AuthProvider", "Firebase Auth", "Firestore / Router"],
        [
            (0, 1, "Nhập email, mật khẩu hoặc tên"),
            (1, 2, "Gọi login() hoặc register()"),
            (2, 3, "signInWithEmailAndPassword() / createUserWithEmailAndPassword()"),
            (3, 2, "Trả Firebase user"),
            (2, 4, "setDoc(users/{uid}) nếu đăng ký"),
            (3, 2, "onAuthStateChanged(firebaseUser)"),
            (2, 4, "getDoc(users/{uid}) và router.replace('/(tabs)')"),
        ],
    )


def generate_google_flow_png(output_path: Path) -> None:
    draw_lifeline_diagram(
        output_path,
        "Luồng đăng nhập Google",
        ["Người dùng", "UI", "Google SDK", "AuthProvider", "Firebase / Firestore"],
        [
            (0, 1, "Chọn tiếp tục với Google"),
            (1, 3, "loginWithGoogle()"),
            (3, 2, "hasPlayServices(), signIn()"),
            (2, 3, "Trả idToken"),
            (3, 4, "GoogleAuthProvider.credential(idToken)"),
            (3, 4, "signInWithCredential(auth, credential)"),
            (3, 4, "getDoc(users/{uid})"),
            (3, 4, "setDoc(users/{uid}) nếu chưa có hồ sơ"),
        ],
    )


def generate_profile_flow_png(output_path: Path) -> None:
    draw_lifeline_diagram(
        output_path,
        "Luồng cập nhật hồ sơ cá nhân",
        ["Người dùng", "ProfileModal", "User Service", "Image Service", "Cloudinary", "Firestore / AuthProvider"],
        [
            (0, 1, "Chọn ảnh và nhập tên hiển thị"),
            (1, 2, "updateUser(uid, userData)"),
            (2, 3, "uploadFileToCloudinary(image, 'users') nếu có ảnh mới"),
            (3, 4, "POST multipart/form-data"),
            (4, 3, "Trả secure_url"),
            (2, 5, "updateDoc(users/{uid}, userData)"),
            (1, 5, "updateUserData(uid) để đồng bộ context"),
        ],
    )


def generate_erd_png(output_path: Path) -> None:
    if Image is None:
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1600, 1100), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_font = load_font(36)
    header_font = load_font(23)
    field_font = load_font(18)
    relation_font = load_font(18)

    draw.text((800, 34), "Mô hình dữ liệu Firestore liên quan", fill="#0f172a", font=title_font, anchor="mt")

    entities = {
        "USERS": ((90, 125, 430, 365), ["uid PK", "name", "email", "image", "createdAt"]),
        "WALLETS": ((610, 125, 990, 390), ["id PK", "uid FK", "name", "amount", "totalIncome", "totalExpenses", "image", "created"]),
        "TRANSACTIONS": ((1110, 125, 1510, 430), ["id PK", "uid FK", "walletId FK", "type", "amount", "category", "date", "description", "image"]),
        "BUDGET": ((610, 585, 990, 805), ["id PK", "walletId FK", "type", "amount"]),
        "NOTIFICATIONS": ((90, 585, 430, 805), ["id PK", "uid FK", "title", "description", "type", "created"]),
    }

    for name, (xy, fields) in entities.items():
        x1, y1, x2, y2 = xy
        draw.rounded_rectangle(xy, radius=16, fill="#ede9fe", outline="#334155", width=2)
        draw.rounded_rectangle((x1, y1, x2, y1 + 50), radius=16, fill="#ddd6fe", outline="#334155", width=2)
        draw.text(((x1 + x2) / 2, y1 + 25), name, fill="#0f172a", font=header_font, anchor="mm")
        y = y1 + 72
        for field in fields:
            draw.text((x1 + 26, y), field, fill="#334155", font=field_font)
            y += 28

    relation_lines = [
        ([(430, 230), (610, 230)], "owns", (520, 208)),
        ([(990, 230), (1110, 230)], "contains", (1050, 208)),
        ([(430, 145), (520, 88), (1210, 88), (1210, 125)], "creates", (820, 70)),
        ([(260, 365), (260, 585)], "receives", (260, 475)),
        ([(800, 390), (800, 585)], "has", (830, 485)),
    ]

    for points, label, label_pos in relation_lines:
        poly_arrow(draw, points, fill="#334155", width=3)
        bbox = draw.textbbox(label_pos, label, font=relation_font, anchor="mm")
        draw.rounded_rectangle(
            (bbox[0] - 8, bbox[1] - 4, bbox[2] + 8, bbox[3] + 4),
            radius=6,
            fill="#f8fafc",
        )
        draw.text(label_pos, label, fill="#0f172a", font=relation_font, anchor="mm")

    note = "Firestore là NoSQL; mỗi collection gồm nhiều document. UID từ Firebase Auth là khóa liên kết chính cho dữ liệu người dùng."
    draw.rounded_rectangle((120, 930, 1480, 1010), radius=14, fill="#ffffff", outline="#cbd5e1", width=2)
    draw.text((800, 970), note, fill="#334155", font=field_font, anchor="mm")
    img.save(output_path)


def generate_architecture_png(output_path: Path) -> None:
    if Image is None:
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1600, 980), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_font = load_font(38)
    lane_font = load_font(24)
    note_font = load_font(19)

    draw.text(
        (800, 38),
        "Kiến trúc Auth - Profile - Firebase - Ngôn ngữ",
        fill="#0f172a",
        font=title_font,
        anchor="mt",
    )

    lanes = [
        ((70, 120, 1530, 300), "1. Tầng giao diện", "#eff6ff"),
        ((70, 335, 1530, 535), "2. Tầng xử lý trong app", "#f0fdf4"),
        ((70, 570, 1530, 755), "3. Cấu hình và dịch vụ ngoài", "#fffbeb"),
    ]
    for xy, label, fill in lanes:
        draw.rounded_rectangle(xy, radius=20, fill=fill, outline="#cbd5e1", width=2)
        draw.text((xy[0] + 24, xy[1] + 18), label, fill="#0f172a", font=lane_font)

    draw_box(draw, (120, 170, 400, 275), "Auth Screens", ["Login, Register", "Forgot / Change password"], "#e0f2fe")
    draw_box(draw, (520, 170, 800, 275), "Profile Screens", ["Profile, Settings", "ProfileModal"], "#e0f2fe")
    draw_box(draw, (1245, 170, 1490, 275), "Root Layout", ["AuthProvider", "i18n init"], "#e0f2fe")

    draw_box(draw, (120, 390, 400, 505), "AuthProvider", ["login, register, Google", "resetPassword, changePassword"], "#dcfce7")
    draw_box(draw, (520, 390, 800, 505), "User Service", ["updateUser(uid, data)", "đồng bộ hồ sơ"], "#dcfce7")
    draw_box(draw, (905, 390, 1185, 505), "Image Service", ["uploadFileToCloudinary()", "getProfileImage()"], "#dcfce7")
    draw_box(draw, (1245, 390, 1490, 505), "i18n Config", ["DEFAULT_LANGUAGE = vi", "translations"], "#dcfce7")

    draw_box(draw, (120, 625, 360, 725), "Firebase Auth", ["Email / Google", "phiên đăng nhập"], "#fee2e2")
    draw_box(draw, (390, 625, 510, 725), "Config", ["Auth", "DB"], "#fef3c7")
    draw_box(draw, (520, 625, 800, 725), "Firestore", ["collection users", "profile data"], "#ede9fe")
    draw_box(draw, (905, 625, 1185, 725), "Cloudinary", ["upload ảnh", "secure_url"], "#fee2e2")
    draw_box(draw, (1245, 625, 1490, 725), "Dịch", ["vi / en"], "#fef3c7")

    # Mỗi cột đi theo một đường chính để sơ đồ dễ đọc và không chồng chéo.
    arrow(draw, (260, 275), (260, 390))
    arrow(draw, (260, 505), (260, 625))

    arrow(draw, (660, 275), (660, 390))
    arrow(draw, (660, 505), (660, 625))

    arrow(draw, (800, 448), (905, 448))
    arrow(draw, (1045, 505), (1045, 625))

    arrow(draw, (1368, 275), (1368, 390))
    arrow(draw, (1368, 505), (1368, 625))

    # Firebase Config cấp cấu hình cho Auth và Firestore bằng hai mũi tên ngắn.
    arrow(draw, (390, 675), (360, 675))
    arrow(draw, (510, 675), (520, 675))

    draw.rounded_rectangle((90, 825, 1510, 925), radius=16, fill="#ffffff", outline="#cbd5e1", width=2)
    wrapped = textwrap.wrap(
        "UI gọi AuthProvider/UserService để xử lý nghiệp vụ. Firebase Auth quản lý phiên đăng nhập, "
        "Firestore lưu hồ sơ mở rộng. Ảnh đại diện upload lên Cloudinary trước khi URL ảnh được lưu vào users. "
        "i18next quản lý key dịch và mặc định dùng tiếng Việt.",
        width=130,
    )
    y = 852
    for line in wrapped:
        draw.text((120, y), line, fill="#334155", font=note_font)
        y += 26

    img.save(output_path)


def generate_mermaid_png(markdown_path: Path, code: str, sequence_number: int) -> Path | None:
    diagrams_dir = markdown_path.parent / "diagrams"
    normalized = code.lower()

    if "flowchart" in normalized and "authprovider" in normalized:
        output_path = diagrams_dir / "auth_profile_architecture.png"
        generate_architecture_png(output_path)
    elif "erdiagram" in normalized:
        output_path = diagrams_dir / "auth_profile_data_model.png"
        generate_erd_png(output_path)
    elif "googlesign" in normalized or "idtoken" in normalized:
        output_path = diagrams_dir / "google_login_flow.png"
        generate_google_flow_png(output_path)
    elif "profilemodal" in normalized and "cloudinary" in normalized:
        output_path = diagrams_dir / "profile_update_flow.png"
        generate_profile_flow_png(output_path)
    elif "onauthstatechanged" in normalized or "login/register" in normalized:
        output_path = diagrams_dir / "auth_register_login_flow.png"
        generate_auth_flow_png(output_path)
    else:
        output_path = diagrams_dir / f"mermaid_diagram_{sequence_number}.png"
        generate_architecture_png(output_path)

    return output_path if output_path.exists() else None


def add_image(doc: Document, markdown_path: Path, image_ref: str, alt_text: str) -> None:
    source = (markdown_path.parent / image_ref).resolve()
    if source.suffix.lower() == ".svg":
        png_path = source.with_suffix(".png")
        generate_architecture_png(png_path)
        source = png_path

    if source.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(source), width=Inches(6.4))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(alt_text or source.stem).italic = True


def add_diagram_image(doc: Document, image_path: Path, caption_text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.4))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run(caption_text).italic = True


def convert_markdown(markdown_path: Path, output_path: Path) -> None:
    doc = Document()
    apply_document_styles(doc)

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    first_heading = True
    mermaid_count = 0

    while index < len(lines):
        line = lines[index].rstrip()

        if not line:
            index += 1
            continue

        if line.strip() == "---PAGEBREAK---":
            doc.add_page_break()
            index += 1
            continue

        if line.startswith("```"):
            language = line.strip("`").strip()
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            if language == "mermaid":
                mermaid_count += 1
                diagram_path = generate_mermaid_png(markdown_path, "\n".join(code_lines), mermaid_count)
                if diagram_path:
                    add_diagram_image(doc, diagram_path, f"Sơ đồ {mermaid_count}")
                continue
            if language:
                add_paragraph(doc, f"Code block: {language}", style="Heading 3")
            for code_line in code_lines:
                paragraph = doc.add_paragraph(style="CodeBlock")
                shade_paragraph(paragraph)
                paragraph.add_run(code_line if code_line else " ")
            continue

        image_match = IMAGE_RE.search(line)
        if image_match:
            add_image(doc, markdown_path, image_match.group(2), image_match.group(1))
            index += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(doc, table_lines)
            continue

        heading_match = HEADING_RE.match(line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            if first_heading and level == 1:
                paragraph = doc.add_paragraph(style="ReportTitle")
                add_formatted_text(paragraph, text)
                first_heading = False
            else:
                style = f"Heading {min(level, 3)}"
                add_paragraph(doc, text, style=style)
            index += 1
            continue

        if line.startswith("- "):
            add_paragraph(doc, line[2:].strip(), style="List Bullet")
            index += 1
            continue

        number_match = NUMBER_RE.match(line)
        if number_match:
            add_paragraph(doc, number_match.group(1).strip(), style="List Number")
            index += 1
            continue

        add_paragraph(doc, line)
        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    convert_markdown(args.input, args.output)


if __name__ == "__main__":
    main()
